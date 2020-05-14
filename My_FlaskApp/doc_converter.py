# -*- coding: utf-8 -*-
"""
Created on Wed May 13 15:15:22 2020

@author: Rajkumar
"""
import os
#import win32com.client as client
import comtypes.client as client
import pythoncom
from docx import Document
from pptx import Presentation
import img2pdf
from PIL import Image
import PyPDF2
import json
from flask import render_template, request

directory = os.getcwd()
path = os.getcwd()
import blob

def converter(raw_file, file_name, raw_format='word'):
    pythoncom.CoInitialize()
    file_path = directory + "/tmp/{}.pdf".format(file_name)
    if raw_format == 'word':
        docx_path = directory + '/tmp/{}.docx'.format(file_name)
        try:
            document = Document(raw_file)
            document.save(docx_path)            
        except PermissionError:
            os.chmod(docx_path, 0o777)
            os.remove(docx_path)
            document = Document(raw_file)
            document.save(docx_path)
#        word = client.DispatchEx("Word.Application")
        word = client.CreateObject("Word.Application")
        doc = word.Documents.Open(docx_path)
        try:
            doc.SaveAs(file_path, FileFormat=17)
            doc.Close()
            word.Quit()
        except Exception as e:
            print(e)
            os.chmod(file_path, 0o777)
            os.remove(file_path)
            doc.SaveAs(file_path, FileFormat=17)
            doc.Close()
            word.Quit()
    elif raw_format == 'ppt':
        file_path = directory + "/tmp//{}.pdf".format(file_name)
        pptx_path = directory + '/tmp/{}.pptx'.format(file_name)
        try:
            document = Presentation(raw_file)
            document.save(pptx_path)            
        except PermissionError:
            os.chmod(pptx_path, 0o777)
            os.remove(pptx_path)
            document = Presentation(raw_file)
            document.save(pptx_path)
#        powerpoint = client.DispatchEx("Powerpoint.Application")
        powerpoint = client.CreateObject("Powerpoint.Application")
        deck = powerpoint.Presentations.Open(pptx_path)
        try:
            deck.SaveAs(file_path, FileFormat=32) # formatType = 32 for ppt to pdf
            deck.Close()
            powerpoint.Quit()
        except Exception as e:
            print(e)
            os.chmod(file_path, 0o777)
            os.remove(file_path)
            deck.SaveAs(file_path, FileFormat=32)
            deck.Close()
            powerpoint.Quit()
    elif raw_format == 'image':
        file_path = directory + "/tmp//{}.pdf".format(file_name)
        extension = os.path.splitext(file_name)[-1].lower()
        image_path = directory + "/tmp//{0}.{1}".format(file_name, extension)
        raw_file.save(image_path)
        try:
            with open(file_path,"wb") as f:
                f.write(img2pdf.convert(image_path))
        except Exception as e:
            print(e)
            image = Image.open(image_path)
#            image.background_color = Color('white')
            image = image.convert("RGB")
            image.alpha_channel = False
            with open(file_path,"wb") as f:
                f.write(img2pdf.convert(image_path))
    return None


def upload_doc():  
    transporter_name = 'a'
    transporters = ['a', 'b']
    popup = 'Profile Successfully Updated'
    popup = json.dumps(popup)
    return render_template('doc_uploader.html', popup=popup,
                           transporters=transporters, transporter_name=transporter_name,
                           userid='ayx')
    

def convert():
    global x
    file = request.files['profile']
    print(file)
    x = file
    extension = os.path.splitext(file.filename)[-1].lower()
    name = os.path.splitext(file.filename)[0].lower()
    if extension == '.pdf':        
        profile = PyPDF2.PdfFileReader(file)
        pages_no = profile.numPages
        output = PyPDF2.PdfFileWriter()
        for i in range(pages_no):
            try:
                profile = PyPDF2.PdfFileReader(file)
                if profile.isEncrypted:
                    profile.decrypt("")
                output.addPage(profile.getPage(i))
            except Exception as e:
                print(str(e))
#        blob.writeBlob_stream(output, 'smple.pdf', container_name='rajkumar')
        with open("tmp/%s.pdf" % file.filename, "wb") as out:
            output.write(out)
    elif extension == '.docx':
        converter(file, name, raw_format='word')
    elif extension == '.pptx':
        converter(file, name, raw_format='ppt')
    elif extension in ['.png', '.jpeg', '.jpg']:
        converter(file, name, raw_format='image')
    else:
        return 'please upload pdf/docx/pptx/image format file'
    filename = name + '.pdf'
    filepath = path + '/tmp/{}.pdf'.format(name)
    blob.writeBlob(filepath, filename, container_name='rajkumar')
    transporter_name = 'a'
    transporters = ['a', 'b']
    popup = 'Profile Successfully Updated'
    popup = json.dumps(popup)
    return render_template('doc_uploader.html', popup=popup,
                           transporters=transporters, transporter_name=transporter_name,
                           userid='ayx')